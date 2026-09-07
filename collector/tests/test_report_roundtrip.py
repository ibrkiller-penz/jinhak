"""
collector/tests/test_report_roundtrip.py — 연세대 예시 라운드트립 무결성 테스트
"""
from pathlib import Path
import openpyxl
import pytest

from collector.scraper.parsers.base import Snapshot, Table
from collector.report.report_builder import write_report_sheet, _norm


def split_example_into_snapshots(ws, last_row: int = 231) -> list[Snapshot]:
    """연세대 예시(정상 구간 1~231행)를 5개 스냅샷으로 분해."""
    labels = [ws.cell(1, c).value for c in (4, 6, 8, 10, 12)]
    snaps = [Snapshot(label=l, captured_at="2026-09-07T20:00:00+09:00", tables=[]) for l in labels]

    # 표 경계 추출: 제목행('현황' 포함, B열 None) / 요약표는 3~4행
    blocks: list[tuple[str | None, int, int]] = [(None, 3, 4)]
    r = 5
    while r <= last_row:
        a = ws.cell(r, 1).value
        if a and "현황" in str(a) and ws.cell(r, 2).value is None:
            title = str(a)
            start = r + 1
            end = start
            while end + 1 <= last_row and not (
                ws.cell(end + 1, 1).value and "현황" in str(ws.cell(end + 1, 1).value) and ws.cell(end + 1, 2).value is None
            ):
                end += 1
            blocks.append((title, start, end))
            r = end + 1
        else:
            r += 1

    for title, start, end in blocks:
        for si in range(5):
            jc, gc = 4 + si * 2, 5 + si * 2
            rows = []
            for rr in range(start, end + 1):
                rows.append([
                    ws.cell(rr, 1).value,
                    ws.cell(rr, 2).value,
                    ws.cell(rr, 3).value,
                    ws.cell(rr, jc).value,
                    ws.cell(rr, gc).value,
                ])
            snaps[si].tables.append(Table(title=title, rows=rows))
    return snaps


def test_yonsei_roundtrip():
    # 프로젝트 루트 기준 샘플 파일 탐색
    root = Path(__file__).resolve().parent.parent.parent
    sample_file = root / "samples" / "취합양식_원본.xlsx"
    assert sample_file.exists(), f"샘플 파일이 없습니다: {sample_file}"

    wb_src = openpyxl.load_workbook(sample_file, data_only=True)
    ws = wb_src["연세대(2025학년도 예시)"]
    last_row = 231  # 정상 구간 검증

    snaps = split_example_into_snapshots(ws, last_row)
    assert len(snaps) == 5
    assert len(snaps[0].tables) > 1

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)
    warns = write_report_sheet(wb_out, "연세대", snaps)

    test_out_dir = root / "out" / "test_results"
    test_out_dir.mkdir(parents=True, exist_ok=True)
    out_path = test_out_dir / "연세대_라운드트립_결과.xlsx"
    wb_out.save(out_path)

    # 생성된 시트와 원본 시트 셀 단위 비교
    ws_gen = wb_out["연세대"]
    diff_count = 0
    diff_details = []

    for r in range(1, last_row + 1):
        for c in range(1, 14):
            val_src = ws.cell(r, c).value
            val_gen = ws_gen.cell(r, c).value
            if _norm(val_src) != _norm(val_gen):
                diff_count += 1
                diff_details.append(f"Row {r}, Col {c}: 원본={val_src!r} vs 생성={val_gen!r}")

    assert diff_count == 0, f"셀 불일치 {diff_count}건 발생:\n" + "\n".join(diff_details[:10])
    assert len(warns) == 0, f"예상치 못한 병합 경고 {len(warns)}건 발생"


if __name__ == "__main__":
    test_yonsei_roundtrip()
    print("✅ 연세대 라운드트립 테스트 100% 통과 (불일치 0건, 경고 0건)")
