"""
collector/storage/backdata_xlsx.py — 시간별 단일 회차 백데이터 엑셀 파일 생성
"""
from __future__ import annotations
from pathlib import Path
import openpyxl
from collector.scraper.parsers.base import Snapshot, Table


def save_backdata_xlsx(
    output_path: Path,
    univ_name: str,
    ratio_url: str | None,
    snapshot: Snapshot,
    screenshot_name: str | None = None,
) -> Path:
    """
    현행 매크로 산출물을 대체하는 단일 회차 백데이터 xlsx 생성.
    - 시트명: snapshot.label (예: "09월07일20시")
    - 내용: tables를 A1부터 순서대로 기록. 표 사이 1행 공백.
    - _meta 시트: 대학명, URL, 수집시각, 캡쳐파일명 등 메타데이터.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    
    # 1. 메인 데이터 시트 (라벨명)
    ws = wb.active
    ws.title = snapshot.label[:31]

    current_row = 1
    for t in snapshot.tables:
        if t.title:
            ws.cell(current_row, 1, t.title)
            current_row += 1

        for row in t.rows:
            for col_idx, val in enumerate(row, start=1):
                # 숫자 변환 시도
                cell_val = val
                if isinstance(val, str):
                    s = val.strip()
                    if s.isdigit():
                        cell_val = int(s)
                ws.cell(current_row, col_idx, cell_val)
            current_row += 1

        # 표 간격 1행 공백
        current_row += 1

    # 2. _meta 시트
    ws_meta = wb.create_sheet("_meta")
    ws_meta.append(["항목", "값"])
    ws_meta.append(["대학명", univ_name])
    ws_meta.append(["경쟁률 URL", ratio_url or ""])
    ws_meta.append(["수집 회차 라벨", snapshot.label])
    ws_meta.append(["수집 시각", snapshot.captured_at])
    ws_meta.append(["수집 상태", snapshot.status])
    ws_meta.append(["스크린샷 파일명", screenshot_name or ""])
    ws_meta.append(["공지/안내문", snapshot.notice or ""])
    ws_meta.append(["총모집인원", snapshot.summary.get("mojip", 0)])
    ws_meta.append(["총지원인원", snapshot.summary.get("jiwon", 0)])
    ws_meta.append(["총경쟁률", snapshot.summary.get("ratio", "-")])

    wb.save(output_path)
    return output_path
